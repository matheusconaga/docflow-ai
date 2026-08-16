from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel
from typing import List

from app.db.database import get_db
from app.core.security import get_current_user
from app.models.user import User
from app.models.lesson_plan_template import LessonPlanTemplate

router = APIRouter(prefix="/templates", tags=["templates"])

class TemplateResponse(BaseModel):
    id: str
    name: str
    html_content: str
    created_at: str

class TemplateCreateRequest(BaseModel):
    name: str
    html_content: str

@router.get("/download-base")
def download_base_template(current_user: User = Depends(get_current_user)):
    """
    Gera um DOCX 'esqueleto' básico para o professor começar a editar.
    """
    import io
    from docx import Document
    
    doc = Document()
    doc.add_heading('Plano de Aula - Seu Logo Aqui', 0)
    
    doc.add_heading('Informações Básicas', level=1)
    doc.add_paragraph('Disciplina: {{ subject }}')
    doc.add_paragraph('Assunto: {{ topic }}')
    doc.add_paragraph('Turma: {{ class_name }}')
    doc.add_paragraph('Duração: {{ duration }}')
    
    doc.add_heading('Objetivos e Estratégias', level=1)
    doc.add_paragraph('Objetivos de Aprendizagem:')
    doc.add_paragraph('{{ objectives }}')
    
    doc.add_paragraph('Metodologias:')
    doc.add_paragraph('{{ methodologies }}')
    
    doc.add_paragraph('Recursos Necessários:')
    doc.add_paragraph('{{ resources }}')
    
    doc.add_heading('Estrutura da Aula', level=1)
    doc.add_paragraph('{{ content }}')
    
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": 'attachment; filename="modelo_base_educassist.docx"'
        }
    )

@router.get("/", response_model=List[TemplateResponse])
def get_templates(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    templates = db.query(LessonPlanTemplate).filter(
        LessonPlanTemplate.teacher_id == current_user.id
    ).order_by(LessonPlanTemplate.created_at.desc()).all()
    
    return [
        TemplateResponse(
            id=t.id,
            name=t.name,
            html_content=t.html_content,
            created_at=t.created_at.isoformat()
        ) for t in templates
    ]

@router.post("/", response_model=TemplateResponse)
async def create_template(
    data: TemplateCreateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    template = LessonPlanTemplate(
        teacher_id=current_user.id,
        name=data.name,
        html_content=data.html_content
    )
    db.add(template)
    db.commit()
    db.refresh(template)
    
    return TemplateResponse(
        id=template.id,
        name=template.name,
        html_content=template.html_content,
        created_at=template.created_at.isoformat()
    )

@router.delete("/{template_id}")
def delete_template(
    template_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    template = db.query(LessonPlanTemplate).filter(
        LessonPlanTemplate.id == template_id,
        LessonPlanTemplate.teacher_id == current_user.id
    ).first()
    
    if not template:
        raise HTTPException(status_code=404, detail="Template não encontrado")
        
    db.delete(template)
    db.commit()
    
    return {"message": "Template deletado com sucesso"}
